"""D176 — màn ISO → Xuất báo cáo: hồ sơ theo dõi cho đoàn kiểm tra ra Excel / PDF.

Vì sao phải có bài này:
  · Đoàn đang ngồi chờ: một biểu mẫu in lỗi không được làm hỏng cả tệp; kỳ trống phải nói "không có bản ghi".
  · Excel phải là ĐÚNG tờ in (ô gộp, tiêu đề, số) — không phải một bảng thô lệch với tờ giấy; chữ "=…" người gõ không
    được biến thành công thức chạy trong Excel của đoàn.
  · PDF: tờ A4 ngang không được dựng dọc (mất cột); bìa đánh đúng số trang; bookmark đúng biểu mẫu.
  · Job nền: job trùng / CHẠY NGAY không dựng hai lần; lỗi phải nằm trên lần xuất, không chỉ trong Error Log.
  · DocType khớp API (ô, lựa chọn Select) — Frappe kiểm Select lúc lưu (bài học D157).

Nạp sx/qc/xuat_bao_cao.py, sx/api/qc_xuatbc.py, qc_hoso và các hàm in THẬT; frappe giả; openpyxl thật (Frappe có sẵn
trên site — máy test: pip install --user openpyxl).
Chạy: python3 scripts/test-xuatbc.py   (verify.sh gọi sẵn)
"""

import io
import json
import os
import shutil
import sys
import tempfile
import types
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

try:
    import openpyxl
except ImportError:
    print("Thiếu openpyxl — pip install --user openpyxl (trên site Frappe có sẵn).\n\nXUATBC-FAIL (thiếu openpyxl)")
    sys.exit(1)

import fakefrappe as F  # noqa: E402

FR = F.cai()
Q = F.nap_qc()
for t in ("khieu_nai", "ncc", "kiem_xe", "tiep_nhan", "rework", "attp", "ho_so", "xuat_bao_cao"):
    F.nap(f"sx.qc.{t}", f"sx/qc/{t}.py")
for t in ("qc_ncc", "qc_attp", "qc_cat", "qc_dvgh", "qc_khieunai", "qc_kiemnghiem", "qc_kiemxe", "qc_rework",
          "qc_thietbi", "qc_tiepnhan", "qc_hoso", "qc_xuatbc"):
    F.nap(f"sx.api.{t}", f"sx/api/{t}.py")
HS = sys.modules["sx.qc.ho_so"]
XB = sys.modules["sx.qc.xuat_bao_cao"]
HOSO = sys.modules["sx.api.qc_hoso"]
CAT = sys.modules["sx.api.qc_cat"]
API = sys.modules["sx.api.qc_xuatbc"]
CT = F.nap("xbc_ctl", "sx/qc/doctype/sx_xuat_bao_cao/sx_xuat_bao_cao.py")
F.dang_ky(XB.PT, CT.SXXuatBaoCao)
CT_HS = F.nap("hs_ctl", "sx/qc/doctype/sx_ho_so_danh_muc/sx_ho_so_danh_muc.py")
F.dang_ky(HS.PT, CT_HS.SXHoSoDanhMuc)
P149 = F.nap("sx.patches.d149_ho_so", "sx/patches/d149_ho_so.py")
FR.local = types.SimpleNamespace(response=F.Doc())
JOB, LOG = [], []
FR.enqueue = lambda method, **k: JOB.append((method, k))
FR.log_error = lambda **k: LOG.append(k)
FR.get_traceback = lambda: "traceback giả"
FR.db.sql = lambda *a, **k: []
FR.db.rollback = lambda: None
kiem, thu = F.kiem, F.thu
F.dat_ngay("2026-10-09")

KHO = tempfile.mkdtemp(prefix="sx-xuatbc-")


class Tep(F.DocThuong):
    """File giả: ghi nội dung ra đĩa (private/files), đọc lại qua get_full_path — như test-hoso."""

    def insert(self, ignore_permissions=False, **k):
        with open(os.path.join(KHO, self["file_name"]), "wb") as fh:
            fh.write(self.pop("content"))
        super().insert()
        self["file_url"] = f"/private/files/{self['file_name']}"
        F.bang("File")[self["name"]]["file_url"] = self["file_url"]
        return self

    def get_full_path(self):
        return os.path.join(KHO, self["file_name"])

    def _goi(self, ten):            # delete_doc giả gọi on_trash / after_delete
        pass


F.LOP["File"] = Tep


def lan(name):
    return F.bang(XB.PT)[name]


# ═══ 1. Hàm thuần: tờ in → khối ══════════════════════════════════════════════
print("-- tờ in (HTML) → khối: dòng chữ, bảng --")
H = """<!doctype html><html><head><meta charset="utf-8"><title>bỏ</title><style>@page { size: A4 landscape;
  margin: 10mm 8mm; } .x { color: red }</style></head><body>
<table class="sx-dau-trang"><tr><td>CÔNG TY</td><td><div>NHẬT KÝ CÁT</div><div>tháng 9</div></td>
  <td><b>BM.08.03</b><br>Lần BH: 01</td></tr></table>
<h1>Tiêu   đề</h1><div class="ky">Ghi chú:<br>· Sáng: ổn
     dòng tiếp &amp; hết</div>
<table><tr><th rowspan="2">#</th><th colspan="2">Nhiệt độ</th><th>KL</th></tr>
 <tr><th>Sáng</th><th>Chiều</th><th>x</th></tr>
 <tr><td>1</td><td>262</td><td>99.5</td><td>1.250</td></tr>
 <tr><td colspan="99">Tổng</td></tr>
 <tr><td>2</td><td><table><tr><td>a</td><td>b</td></tr><tr><td>c</td></tr></table></td><td>=1+1</td><td>09</td></tr>
</table><p>Người lập: A</p></body></html>"""
K = XB.tach(H)
dong = [(k[1], k[2]) for k in K if k[0] == "dong"]
bang = [k[1] for k in K if k[0] == "bang"]
kiem("bỏ <style>, <title>; chữ ngoài bảng mỗi dòng một hàng (theo <br>, khối); khoảng trắng mã nguồn gộp; &amp; → &",
     dong == [("Tiêu đề", "tieu_de"), ("Ghi chú:", ""), ("· Sáng: ổn dòng tiếp & hết", ""), ("Người lập: A", "")],
     dong)
kiem("đầu trang chung (bảng 3 ô) giữ xuống dòng trong ô (<div>, <br>)",
     [o["text"] for o in bang[0][0]] == ["CÔNG TY", "NHẬT KÝ CÁT\ntháng 9", "BM.08.03\nLần BH: 01"], bang[0])
kiem("ô tiêu đề <th> đánh dấu; colspan / rowspan giữ nguyên", [(o["text"], o["cs"], o["rs"], o["th"]) for o in
                                                               bang[1][0]] == [("#", 1, 2, True), ("Nhiệt độ", 2, 1, True),
                                                                               ("KL", 1, 1, True)])
kiem("bảng lồng trong ô → chữ của ô (hàng một dòng, ô cách ' | ')", bang[1][4][1]["text"] == "a | b\nc", bang[1][4])
o_ds, cot = XB.luoi(bang[1])
vi_tri = {o["text"]: (r, c, o["cs"], o["rs"]) for r, c, o in o_ds}
kiem("lưới như trình duyệt: ô dưới rowspan dồn sang phải; colspan quá bề ngang bảng (99) cắt ở mép (4 cột)",
     cot == 4 and vi_tri["Sáng"] == (1, 1, 1, 1) and vi_tri["x"] == (1, 3, 1, 1) and vi_tri["Tổng"] == (3, 0, 4, 1),
     (cot, vi_tri))
kiem("rowspan vượt đáy bảng → cắt ở đáy", XB.luoi([[{"text": "a", "cs": 1, "rs": 9, "th": False}]])[0][0][2]["rs"] == 1)
kiem("tờ thiếu </table> vẫn lấy được bảng", XB.tach("<table><tr><td>a</td><td>b")[0][1][0][1]["text"] == "b")
kiem("ô chắc là số → số (262, 99.5, 0, -3); '1.250' (có thể là nghìn), '09' (mã), '99.5%', '1e5' giữ chữ",
     [XB.so(t) for t in ("262", "99.5", "0", "-3", "1.250", "09", "99.5%", "1e5", "")] == [262, 99.5, 0, -3, "1.250", "09",
                                                                                          "99.5%", "1e5", ""])
da = set()
ten = [XB.ten_sheet(t, da) for t in ("BM.08.01 2026-09", "bm.08.01 2026-09", "a:b/c?d*e[f]g", "x" * 40, "x" * 40,
                                     "'Mục lục'")]
kiem("tên sheet: bỏ ký tự cấm, ≤ 31 ký tự, trùng (kể cả khác hoa thường) thêm (2), (3)",
     ten == ["BM.08.01 2026-09", "bm.08.01 2026-09 (2)", "a-b-c-d-e-f-g", "x" * 31, "x" * 27 + " (2)", "Mục lục"]
     and all(len(t) <= 31 for t in ten), ten)
kiem("nhãn tờ: biểu mẫu một tờ → mã; nhiều tờ → mã + tháng", (XB.nhan_to("BM.08.03", "2026-09.html", False),
                                                               XB.nhan_to("BM.08.02", "so-su-co.csv", True))
     == ("BM.08.03 2026-09", "BM.08.02"))
kg = [XB.kho_giay(f"<style>{c}</style>") for c in ("@page { size: A4 landscape; margin: 10mm 8mm; }",
                                                   "@page{size:A4 portrait;margin:12mm}", "@page { margin: 1cm 2cm 3cm }",
                                                   "", "@page { size: A4; margin: auto }")]
kiem("khổ giấy theo @page: ngang / dọc, lề CSS 1–3 giá trị; không khai / lề lạ → A4 dọc lề 10 mm",
     (kg[0]["orientation"], kg[0]["margin-top"], kg[0]["margin-left"]) == ("Landscape", "10mm", "8mm")
     and (kg[1]["orientation"], kg[1]["margin-right"]) == ("Portrait", "12mm")
     and (kg[2]["margin-top"], kg[2]["margin-right"], kg[2]["margin-bottom"], kg[2]["margin-left"])
     == ("1cm", "2cm", "3cm", "2cm") and kg[3] == kg[4] and kg[3]["margin-top"] == "10mm", kg)
ngang, doc = "<style>@page{size:A4 landscape}</style>a", "<style>@page{size:A4 portrait}</style>b"
kiem("tờ liền nhau cùng khổ gom một lượt dựng; đổi khổ thì lượt mới",
     [(k["orientation"], len(ds)) for k, ds in XB.nhom_theo_kho([ngang, ngang, doc, ngang])]
     == [("Landscape", 2), ("Portrait", 1), ("Landscape", 1)])
g = XB.ghep_html(["<html><body><p>một</p></body></html>", "<p>hai</p>"], "T")
kiem("ghép tờ: lấy phần trong <body>, mỗi tờ sang trang mới, tự khai bảng mã", '<meta charset="utf-8">' in g
     and "<p>một</p>" + XB.NGAT + "<p>hai</p>" in g and g.count("<body>") == 1)
c = XB.tach_csv("﻿Số,Mục,Nội dung\nSC-1,B2,\"a, b\"\n\n")[0][1]
kiem("CSV (sổ sự cố): bỏ BOM, hàng đầu là tiêu đề, ô có dấu phẩy giữ nguyên, bỏ hàng trống",
     [[o["text"] for o in h] for h in c] == [["Số", "Mục", "Nội dung"], ["SC-1", "B2", "a, b"]] and c[0][0]["th"]
     and not c[1][0]["th"])
nb = XB.nhom_bieu_mau({"BM.08.03": "Cát", "BM.08.01": "Vòng", "BM.01.11": "Họp ISO", "XX": "Lạ"},
                      {"BM.08.03": "Kiểm soát sản xuất", "BM.08.01": "Kiểm soát sản xuất", "BM.01.11": "Hệ thống quản lý"},
                      HS.NHOM, bien_ban=["BM.01.11"])
kiem("danh sách chọn theo nhóm danh mục hồ sơ (thứ tự BM.01.04), mã trong nhóm theo thứ tự; chưa có trong danh mục → "
     "nhóm cuối; mỗi biểu mẫu nói sẽ ra tờ nào",
     [(n["ten"], [b["ma"] for b in n["bm"]]) for n in nb] == [("Hệ thống quản lý", ["BM.01.11"]),
                                                              ("Kiểm soát sản xuất", ["BM.08.01", "BM.08.03"]),
                                                              (XB.NGOAI_DANH_MUC, ["XX"])]
     and nb[1]["bm"][0]["ky"] == "mỗi tháng một tờ (tờ từng ngày)" and nb[0]["bm"][0]["ky"] == XB.KY_BIEN_BAN
     and nb[2]["bm"][0]["ky"] == XB.KY_KHAC, nb)
kiem("tên tệp tải về theo kỳ, đúng đuôi", (XB.ten_tep("Excel", "2026-09-01", "2026-10-09"), XB.ten_tep("PDF", "a", "b"))
     == ("ho-so-theo-doi_2026-09-01_2026-10-09.xlsx", "ho-so-theo-doi_a_b.pdf"))

# ═══ 2. Dựng workbook (openpyxl thật) ═════════════════════════════════════════
print("\n-- dựng Excel từ khối --")
bm = [{"ma": "BM.08.03", "ten": "Nhật ký cát rang", "to": [("2026-09.html", H), ("2026-10.html", H)], "loi": ""},
      {"ma": "BM.08.02", "ten": "Sổ sự cố", "to": [("so-su-co.csv", "﻿Số,Phút\nSC-1,25\n")], "loi": ""},
      {"ma": "BM.08.04", "ten": "Xuất xưởng", "to": [], "loi": ""},
      {"ma": "BM.08.01", "ten": "Vòng kiểm", "to": [], "loi": "division by zero"}]
wb = openpyxl.load_workbook(io.BytesIO(XB.dung_excel(openpyxl, {"tieu_de": "HỒ SƠ", "dong": ["Kỳ: x"]}, bm)))
kiem("sheet: Mục lục đầu tiên, mỗi tờ một sheet (nhiều tờ → mã + tháng; một tờ → mã)",
     wb.sheetnames == ["Mục lục", "BM.08.03 2026-09", "BM.08.03 2026-10", "BM.08.02"], wb.sheetnames)
ws = wb["BM.08.03 2026-09"]
gia_tri = {o.coordinate: o.value for h in ws.iter_rows() for o in h if o.value is not None}
kiem("tiêu đề sheet = mã — tên · tháng; khối giữ thứ tự tờ in (đầu trang, tiêu đề, ghi chú, bảng, người lập)",
     gia_tri["A1"] == "BM.08.03 — Nhật ký cát rang · 2026-09" and gia_tri["A3"] == "CÔNG TY"
     and gia_tri["B3"] == "NHẬT KÝ CÁT\ntháng 9" and gia_tri["A5"] == "Tiêu đề" and gia_tri["A8"] == "#"
     and gia_tri["A14"] == "Người lập: A", gia_tri)
gop = sorted(str(r) for r in ws.merged_cells.ranges)
kiem("ô gộp đúng như tờ in: '#' 2 hàng, 'Nhiệt độ' 2 cột, 'Tổng' hết bề ngang", gop == ["A11:D11", "A8:A9", "B8:C8"],
     gop)
kiem("số thành số (262, 99.5), '1.250' / '09' giữ chữ", (ws["B10"].value, ws["C10"].value, ws["D10"].value,
                                                         ws["D12"].value) == (262, 99.5, "1.250", "09"))
kiem("'=1+1' người gõ là CHỮ, không thành công thức", ws["C12"].value == "=1+1" and ws["C12"].data_type == "s")
kiem("<th> đậm nền xám; ô bảng có viền, chữ xuống dòng", ws["B8"].font.b and ws["B8"].fill.fgColor.rgb.endswith("EEEEEE")
     and ws["B10"].border.left.style == "thin" and ws["C9"].border.top.style == "thin" and ws["B10"].alignment.wrap_text)
kiem("thiết lập in theo tờ: A4 ngang, vừa bề ngang trang", ws.page_setup.orientation == "landscape"
     and ws.page_setup.fitToWidth == 1 and ws.sheet_properties.pageSetUpPr.fitToPage)
ml = wb["Mục lục"]
hang = [[o.value for o in h] for h in ml.iter_rows(min_row=4, max_row=9)]
kiem("Mục lục: biểu mẫu → các tờ; kỳ trống nói 'Kỳ này không có bản ghi'; in lỗi nói lỗi",
     hang[:5] == [["STT", "Mã", "Biểu mẫu", "Tờ (bấm để mở)", "Ghi chú"],
                  [1, "BM.08.03", "Nhật ký cát rang", "BM.08.03 2026-09", None],
                  [None, None, None, "BM.08.03 2026-10", None], [2, "BM.08.02", "Sổ sự cố", "BM.08.02", None],
                  [3, "BM.08.04", "Xuất xưởng", None, "Kỳ này không có bản ghi"]]
     and hang[5] == [4, "BM.08.01", "Vòng kiểm", None, "Không in được: division by zero"], hang)
kiem("bấm tên tờ trong Mục lục sang đúng sheet (liên kết trong tệp, không phải đường dẫn ra ngoài)",
     ml["D5"].hyperlink.location == "'BM.08.03 2026-09'!A1" and not ml["D5"].hyperlink.target)
kiem("CSV thành bảng: tiêu đề đậm, số thành số", wb["BM.08.02"]["A3"].value == "Số" and wb["BM.08.02"]["A3"].font.b
     and wb["BM.08.02"]["B4"].value == 25)

# ═══ 3. DocType khớp API ═════════════════════════════════════════════════════
print("\n-- DocType SX Xuat Bao Cao --")
dt = json.load(open("sx/qc/doctype/sx_xuat_bao_cao/sx_xuat_bao_cao.json", encoding="utf-8"))
o = {f["fieldname"]: f for f in dt["fields"]}
kiem("mọi ô API đọc / ghi có trong DocType", set(API.TRUONG) - {"name", "owner", "creation"} <= set(o),
     sorted(set(API.TRUONG) - {"name", "owner", "creation"} - set(o)))
kiem("Select trạng thái / kiểu ĐÚNG hằng số của luật (lệch là Frappe chặn lúc lưu)",
     o["trang_thai"]["options"].split("\n") == list(XB.TRANG_THAI) and o["kieu"]["options"].split("\n") == list(XB.KIEU))
kiem("quyền Desk: Trưởng Ban ISO, quản lý, System Manager", {p["role"] for p in dt["permissions"]}
     == {"ISO Manager", "SX Quan Ly", "System Manager"})
kiem("controller: kiểu lạ / kỳ ngược → chặn", "Kiểu tệp" in (thu(lambda: F.get_doc(
    {"doctype": XB.PT, "kieu": "Word", "tu": "2026-09-01", "den": "2026-09-30"}).insert()) or "")
     and "từ ngày" in (thu(lambda: F.get_doc({"doctype": XB.PT, "kieu": "Excel", "tu": "2026-09-30",
                                             "den": "2026-09-01"}).insert()) or ""))
F.bang(XB.PT).clear()

# ═══ 4. API: chọn, xuất, job, tải ════════════════════════════════════════════
print("\n-- API: danh sách chọn, xuất Excel (job nền) --")
P149.execute()
base = {"docstatus": 1, "ghi_muon": 0, "nhap_lai_tu_giay": 0, "qc_user": "qc@x", "co_san_xuat_bot": 0,
        "reviewed_on": None, "phien_ban": 2}
F.bang("SX QC Round")["QC-1"] = dict(base, name="QC-1", ngay="2026-09-15", luot="Đầu sáng", rang_nhiet_do=262,
                                     started_at=datetime(2026, 9, 15, 7, 0), finished_at=datetime(2026, 9, 15, 7, 30),
                                     creation="2026-09-15 07:00:00")
F.bang("SX Nhat Ky Cat")["CAT-1"] = {"name": "CAT-1", "ngay": "2026-09-15", "so_ngay_dung": 3, "ten_ncc": "Cát Lô",
                                     "thay_cat": 0, "doi_nguon": 0}
F.bang("SX Dien Tap Truy Xuat")["DT-1"] = {"name": "DT-1", "ngay": "2026-09-30", "ket_thuc": "2026-09-30 10:00:00",
                                           "ten_san_pham": "Bánh 1", "lo": "LO-1", "so_phut": 25,
                                           "can_bang_pt": 99.5, "dat": 1}
F.bang("SX Su Co")["SC-1"] = {"name": "SC-1", "ngay": "2026-09-20", "trang_thai": "Mở", "muc_do": "Cao",
                              "xu_ly_ngay": "x", "dien_tap": 0}
F.bang("SX Mau Bien Ban")["BM.99.01"] = {"name": "BM.99.01", "ma": "BM.99.01", "ten": "Biên bản site tự khai", "ngung": 0}
F.bang("SX Mau Bien Ban")["BM.99.02"] = {"name": "BM.99.02", "ma": "BM.99.02", "ten": "Mẫu đã ngừng", "ngung": 1}
F.vai("ISO Manager")
r = API.tong_quan()
nhom = {b["ma"]: n["ten"] for n in r["nhom"] for b in n["bm"]}
kiem("danh sách chọn: mọi biểu mẫu app; nhóm theo danh mục hồ sơ (BM.08.01 → Kiểm soát sản xuất); mẫu biên bản site "
     "tự khai còn dùng có mặt (mẫu đã ngừng thì không)",
     set(HS.BIEU_MAU) <= set(nhom) and nhom["BM.08.01"] == "Kiểm soát sản xuất" and "BM.99.01" in nhom
     and "BM.99.02" not in nhom, {k: nhom.get(k) for k in ("BM.08.01", "BM.99.01")})
kiem("kỳ mặc định như gói zip (3 tháng tới hôm nay); hai kiểu Excel / PDF; chưa có lần xuất",
     (r["tu"], r["den"], r["kieu"], r["ds"]) == ("2026-08-01", "2026-10-09", ["Excel", "PDF"], []))
CHON = ["BM.08.01", "BM.08.03", "BM.08.02", "BM.02.04", "BM.08.04", "TU_CONG_BO", "BM.01.04", "BM.08.01"]
for p, chu in (({"kieu": "Word"}, "Excel hay PDF"), ({"kieu": "Excel"}, "Chọn kỳ"),
               ({"kieu": "Excel", "tu": "2026-10-09", "den": "2026-09-01"}, "trước"),
               ({"kieu": "Excel", "tu": "2024-01-01", "den": "2026-10-09"}, "tối đa 24 tháng — chia làm nhiều lần xuất"),
               ({"kieu": "Excel", "tu": "2026-09-01", "den": "2026-10-09", "bieu_mau": [" "]}, "ít nhất một"),
               ({"kieu": "Excel", "tu": "2026-09-01", "den": "2026-10-09", "bieu_mau": ["BM.08.01", "BM.77.77"]},
                "BM.77.77")):
    kiem(f"xuất: chặn khi {chu}", chu in (thu(lambda p=p: API.xuat(json.dumps(p))) or ""), thu(lambda p=p: API.xuat(
        json.dumps(p))))
kiem("không có lần xuất nào lọt qua khi bị chặn", not F.bang(XB.PT) and not JOB)
x = API.xuat(json.dumps({"kieu": "Excel", "tu": "2026-09-01", "den": "2026-10-09", "bieu_mau": CHON,
                         "ghi_chu": "  Đoàn   Orion  " + "x" * 200}))
d = lan(x["name"])
kiem("ghi lần xuất Đang chờ: biểu mẫu bỏ trùng giữ thứ tự chọn, ghi chú gọn ≤ 140 ký tự",
     d["trang_thai"] == XB.CHO and json.loads(d["bieu_mau"]) == CHON[:-1] and d["so_bieu_mau"] == 7
     and d["ghi_chu"].startswith("Đoàn Orion x") and len(d["ghi_chu"]) == 140, d)
kiem("xếp job nền SAU commit, hàng long, chống trùng theo lần xuất", JOB and JOB[-1][0] == "sx.api.qc_xuatbc.chay"
     and JOB[-1][1]["name"] == x["name"] and JOB[-1][1]["enqueue_after_commit"] and JOB[-1][1]["deduplicate"]
     and JOB[-1][1]["job_id"] == f"sx-xuatbc-{x['name']}" and JOB[-1][1]["queue"] == "long", JOB[-1:])
kiem("chưa xong thì không tải được", "chưa có tệp" in (thu(lambda: API.tai(x["name"])) or ""))
API.chay(x["name"])
d = lan(x["name"])
kq = {k["ma"]: k for k in json.loads(d["ket_qua"])}
kiem("job: Xong, có tệp riêng tư gắn lần xuất, tên theo kỳ, đếm tờ; kết quả từng biểu mẫu",
     d["trang_thai"] == XB.XONG and d["tep"] == "/private/files/ho-so-theo-doi_2026-09-01_2026-10-09.xlsx"
     and d["ten_tep"].endswith(".xlsx") and d["kich_thuoc"] > 1000 and d["bat_dau_luc"] and d["xong_luc"]
     and kq["BM.08.03"]["so_to"] == 2 and kq["BM.08.04"]["so_to"] == 0 and not kq["BM.08.01"]["loi"]
     and next(f for f in F.bang("File").values() if f.get("attached_to_name") == x["name"])["is_private"] == 1,
     (d["trang_thai"], d.get("loi"), kq))
FR.local.response = F.Doc()
API.tai(x["name"])
rs = FR.local.response
wb = openpyxl.load_workbook(io.BytesIO(rs.filecontent))
kiem("tải: kiểu download, tên tệp theo kỳ", rs.type == "download" and rs.filename.endswith("2026-10-09.xlsx"))
kiem("tệp thật: Mục lục + tờ ngày BM.08.01 tháng 9, cát tháng 9 + 10, sổ sự cố, diễn tập, tự công bố, BM.01.04",
     wb.sheetnames == ["Mục lục", "BM.08.01", "BM.08.03 2026-09", "BM.08.03 2026-10", "BM.08.02", "BM.02.04",
                       "TU_CONG_BO", "BM.01.04"], wb.sheetnames)
chu = {s: " ".join(str(o.value) for h in wb[s].iter_rows() for o in h if o.value is not None) for s in wb.sheetnames}
kiem("Mục lục ghi kỳ, người xuất, ghi chú đoàn; BM.08.04 kỳ này không có bản ghi",
     "Kỳ: 01/09/2026 – 09/10/2026" in chu["Mục lục"] and "iso@x" in chu["Mục lục"] and "Đoàn Orion" in chu["Mục lục"]
     and "BM.08.04 Kiểm tra xuất xưởng theo lô Kỳ này không có bản ghi" in chu["Mục lục"], chu["Mục lục"][:600])
kiem("BM.08.01: đúng tờ ngày (đầu trang chung, ngày 15/09/2026, nhiệt độ rang 262 thành số) — có ô gộp hàng bước",
     "Phiếu kiểm tra QC hàng ngày" in chu["BM.08.01"] and "15/09/2026" in chu["BM.08.01"]
     and any(o.value == 262 for h in wb["BM.08.01"].iter_rows() for o in h) and wb["BM.08.01"].merged_cells.ranges,
     chu["BM.08.01"][:400])
kiem("diễn tập: lô, phút (số), cân bằng; tự công bố: bảng sản phẩm", "LO-1" in chu["BM.02.04"]
     and any(o.value == 25 for h in wb["BM.02.04"].iter_rows() for o in h) and "99.5%" in chu["BM.02.04"]
     and "Số bản tự công bố" in chu["TU_CONG_BO"])
kiem("lần xuất hiện ở danh sách gần đây: người xuất, kỳ, có tệp, kết quả", (lambda g: g["name"] == x["name"]
     and g["co_tep"] and g["nguoi"] == "iso@x" and g["tu"] == "2026-09-01" and "tep" not in g
     and len(g["ket_qua"]) == 7)(API.tong_quan()["ds"][0]))
so_tep = lambda n: len([f for f in F.bang("File").values() if f.get("attached_to_name") == n])  # noqa: E731
n_tep = so_tep(x["name"])
kiem("chạy lại job (job trùng) / CHẠY NGAY khi đã xong: không dựng lại, không thêm tệp", (API.chay(x["name"]),
     so_tep(x["name"]))[1] == n_tep == 1 and lan(x["name"])["tep"] == d["tep"]
     and "không còn" in (thu(lambda: API.chay_ngay(x["name"])) or "") and so_tep(x["name"]) == 1)
kiem("trạng thái nhiều lần xuất một lượt (màn hỏi lại)", [y["trang_thai"] for y in API.trang_thai(json.dumps(
    [x["name"]]))] == [XB.XONG] and API.trang_thai("[]") == [])

print("\n-- một biểu mẫu in lỗi, cả lần xuất lỗi, CHẠY NGAY, làm lại, xoá --")
goc = Q.month_sheets
Q.month_sheets = lambda *a: 1 / 0
x2 = API.xuat(json.dumps({"kieu": "Excel", "tu": "2026-09-01", "den": "2026-09-30", "bieu_mau": ["BM.08.01",
                                                                                                    "BM.08.03"]}))
API.chay_ngay(x2["name"])
Q.month_sheets = goc
d2 = lan(x2["name"])
wb2 = openpyxl.load_workbook(io.BytesIO(open(os.path.join(KHO, d2["ten_tep"]), "rb").read()))
kiem("một tờ in lỗi (BM.08.01) → tệp vẫn ra; Mục lục + kết quả nói 'Không in được'; CHẠY NGAY dựng luôn",
     d2["trang_thai"] == XB.XONG and wb2.sheetnames == ["Mục lục", "BM.08.03"]
     and "Không in được: division by zero" in " ".join(str(o.value) for h in wb2["Mục lục"].iter_rows() for o in h)
     and json.loads(d2["ket_qua"])[0]["loi"] == "division by zero", d2)
goc_dung = API._dung
API._dung = lambda doc: 1 / 0
x3 = API.xuat(json.dumps({"kieu": "Excel", "tu": "2026-09-01", "den": "2026-09-30", "bieu_mau": ["BM.08.03"]}))
API.chay(x3["name"])
API._dung = goc_dung
kiem("cả lần xuất hỏng → trạng thái Lỗi kèm lý do trên lần xuất + Error Log", lan(x3["name"])["trang_thai"] == XB.LOI
     and lan(x3["name"])["loi"] == "division by zero" and LOG and "xuất báo cáo" in LOG[-1]["title"])
kiem("làm lại: chỉ lần lỗi; về Đang chờ, xếp job lại", "Chỉ làm lại" in (thu(lambda: API.lam_lai(x["name"])) or "")
     and API.lam_lai(x3["name"])["trang_thai"] == XB.CHO and lan(x3["name"])["loi"] is None
     and JOB[-1][1]["name"] == x3["name"])
API.chay(x3["name"])
kiem("… job chạy lại thành công", lan(x3["name"])["trang_thai"] == XB.XONG)
lan(x3["name"])["trang_thai"] = XB.DANG
kiem("đang dựng thì không xoá", "Đang dựng" in (thu(lambda: API.xoa(x3["name"])) or ""))
lan(x3["name"])["trang_thai"] = XB.XONG
API.xoa(x3["name"])
kiem("xoá lần xuất: mất cả tệp gắn kèm", x3["name"] not in F.bang(XB.PT)
     and not [f for f in F.bang("File").values() if f.get("attached_to_name") == x3["name"]])
for vai, duoc in (("SX QC", False), ("Production Manager", False), ("SX Quan Ly", True)):
    F.vai(vai)
    kiem(f"{vai}: {'xem, xuất, tải được' if duoc else 'không xem, không xuất, không tải, không xoá'}",
         all((thu(f) is None) == duoc for f in (API.tong_quan, lambda: API.tai(x["name"]),
                                                lambda: API.trang_thai(json.dumps([x["name"]])))))
    if not duoc:
        kiem(f"… {vai}: xuất / xoá / chạy ngay bị chặn", all(thu(f) for f in (
            lambda: API.xuat(json.dumps({"kieu": "Excel", "tu": "2026-09-01", "den": "2026-09-30",
                                         "bieu_mau": ["BM.08.03"]})), lambda: API.xoa(x["name"]),
            lambda: API.chay_ngay(x["name"]))))
F.vai("ISO Manager")

# ═══ 5. PDF (get_pdf + PdfWriter giả) ═══════════════════════════════════════
print("\n-- PDF: khổ từng tờ, bìa đánh số trang, bookmark --")
GOI_PDF = []


class GiaWriter:
    def __init__(self):
        self.pages, self.outline = [], []

    def add_page(self, p):
        self.pages.append(p)

    def add_outline_item(self, t, i):
        self.outline.append((t, i))

    def write(self, buf):
        buf.write(b"%PDF-gia " + json.dumps({"pages": self.pages, "outline": self.outline}).encode())


def get_pdf(html, options=None, output=None, smart_shrinking=False):
    GOI_PDF.append((html, dict(options or {}), smart_shrinking))
    if "LOI-PDF" in html:
        raise OSError("wkhtmltopdf: lỗi giả")
    bia = "<h1>" in html and "Trang</th>" in html
    n = (2 if html.count("<tr>") > 7 else 1) if bia else html.count(XB.NGAT) + 1
    for i in range(n):
        output.pages.append(("bìa" if bia else html[html.find("<title>") + 7:html.find("</title>")]) + f"#{i}")
    return output


pdfm = types.ModuleType("frappe.utils.pdf")
pdfm.get_pdf = get_pdf
sys.modules["frappe.utils.pdf"] = pdfm
FR.utils.pdf = pdfm
sys.modules["pypdf"] = types.SimpleNamespace(PdfWriter=GiaWriter)
x4 = API.xuat(json.dumps({"kieu": "PDF", "tu": "2026-09-01", "den": "2026-10-09",
                          "bieu_mau": ["BM.08.01", "BM.08.03", "BM.08.04", "BM.02.04"]}))
API.chay(x4["name"])
d4 = lan(x4["name"])
nd = open(os.path.join(KHO, d4["ten_tep"]), "rb").read()
pdf = json.loads(nd[len(b"%PDF-gia "):])
kiem("PDF Xong, tên .pdf", d4["trang_thai"] == XB.XONG and d4["ten_tep"].endswith(".pdf"), d4)
kho = {h[h.find("<title>") + 7:h.find("</title>")].split(" — ")[0]: (o["orientation"], o["margin-left"], s)
       for h, o, s in GOI_PDF if "<title>" in h and "<title></title>" not in h}
kiem("mỗi biểu mẫu một lượt theo ĐÚNG khổ tờ in: BM.08.01 A4 dọc, BM.08.03 / BM.02.04 A4 ngang lề 8 mm; thu nhỏ bảng "
     "rộng thay vì cắt cột", kho.get("BM.08.01", ("",))[0] == "Portrait" and kho.get("BM.08.03")[:2] == ("Landscape", "8mm")
     and kho.get("BM.02.04")[0] == "Landscape" and all(v[2] for v in kho.values()), kho)
kiem("hai tờ cát (tháng 9, 10) cùng khổ → MỘT lượt dựng, hai trang", sum(1 for h, _o, _s in GOI_PDF
                                                                        if "BM.08.03 —" in h) == 1
     and [p for p in pdf["pages"] if p.startswith("BM.08.03")] == ["BM.08.03 — Nhật ký cát rang#0",
                                                                    "BM.08.03 — Nhật ký cát rang#1"], pdf["pages"])
bia = next(h for h, _o, _s in GOI_PDF if "Trang</th>" in h)
kiem("bìa: trang đầu tệp; mỗi biểu mẫu → trang bắt đầu (bìa 1 trang → BM.08.01 trang 2, cát trang 3, diễn tập 5); "
     "kỳ trống nói rõ", pdf["pages"][0] == "bìa#0" and "<td>BM.08.01</td><td>Vòng kiểm hằng ngày</td><td class='c'>2</td>"
     in bia and "<td class='c'>3</td><td>2 tờ</td>" in bia and "<td class='c'>5</td>" in bia
     and "Kỳ này không có bản ghi" in bia, bia[-900:])
kiem("bookmark: Mục lục, rồi từng biểu mẫu đúng trang (đếm từ 0)", pdf["outline"] == [
    ["Mục lục", 0], ["BM.08.01 — Vòng kiểm hằng ngày", 1], ["BM.08.03 — Nhật ký cát rang", 2],
    ["BM.02.04 — Diễn tập truy xuất (phụ lục)", 4]], pdf["outline"])
GOI_PDF.clear()
CAT.in_bm0803 = (lambda goc_in: (lambda t: goc_in(t) + "LOI-PDF"))(CAT.in_bm0803)
x5 = API.xuat(json.dumps({"kieu": "PDF", "tu": "2026-09-01", "den": "2026-10-09",
                          "bieu_mau": ["BM.08.03", "BM.08.01", "BM.08.04", "BM.02.04", "TU_CONG_BO", "BM.01.04",
                                       "BM.08.02", "BM.01.07"]}))
API.chay(x5["name"])
d5 = lan(x5["name"])
pdf5 = json.loads(open(os.path.join(KHO, d5["ten_tep"]), "rb").read()[len(b"%PDF-gia "):])
bia5 = [h for h, _o, _s in GOI_PDF if "Trang</th>" in h]
kiem("một biểu mẫu dựng PDF hỏng (cát) → tệp vẫn ra, bìa nói 'Không in được: wkhtmltopdf…', không có trang dở",
     d5["trang_thai"] == XB.XONG and "Không in được: wkhtmltopdf: lỗi giả" in bia5[-1]
     and not any(p.startswith("BM.08.03") for p in pdf5["pages"])
     and json.loads(d5["ket_qua"])[0]["loi"] == "wkhtmltopdf: lỗi giả" and json.loads(d5["ket_qua"])[0]["so_to"] == 0,
     d5)
kiem("bìa dài thành 2 trang → dựng lại bìa, số trang các mục dời theo (BM.08.01 trang 3)",
     len(bia5) == 2 and pdf5["pages"][:2] == ["bìa#0", "bìa#1"] and "<td class='c'>3</td>" in bia5[-1]
     and ["BM.08.01 — Vòng kiểm hằng ngày", 2] in pdf5["outline"], (len(bia5), pdf5["outline"]))
kiem("CSV (sổ sự cố) vào PDF thành bảng HTML", any("BM.08.02 —" in h and "<table>" in h for h, _o, _s in GOI_PDF))
x6 = API.xuat(json.dumps({"kieu": "PDF", "tu": "2026-09-01", "den": "2026-10-09", "bieu_mau": ["BM.08.03"]}))
API.chay(x6["name"])
kiem("mọi biểu mẫu đều không dựng được (máy chủ thiếu wkhtmltopdf…) → lần xuất Lỗi, nói lý do",
     lan(x6["name"])["trang_thai"] == XB.LOI and "wkhtmltopdf" in lan(x6["name"])["loi"], lan(x6["name"]))
x7 = API.xuat(json.dumps({"kieu": "PDF", "tu": "2026-09-01", "den": "2026-09-30", "bieu_mau": ["BM.08.04"]}))
API.chay(x7["name"])
kiem("kỳ trống hết (không biểu mẫu nào có tờ) → vẫn ra PDF chỉ có bìa", lan(x7["name"])["trang_thai"] == XB.XONG
     and json.loads(open(os.path.join(KHO, lan(x7["name"])["ten_tep"]), "rb").read()[9:])["pages"] == ["bìa#0"])

shutil.rmtree(KHO, ignore_errors=True)
F.ket_thuc("XUATBC")
